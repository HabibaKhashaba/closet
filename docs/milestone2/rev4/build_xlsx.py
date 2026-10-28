# Traceability matrix v1.5 and API table v1.5: the v1.4 workbooks + the v1.5 rows (same columns).
import sys, copy, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reqs import FRS
from api_v15 import API
SRC, OUT = sys.argv[1], sys.argv[2]
UC = {'FR-81': 'UC-23', 'FR-82': 'UC-23', 'FR-83': 'UC-24', 'FR-84': 'UC-25', 'FR-85': 'UC-27', 'FR-86': 'UC-27', 'FR-87': 'UC-29', 'FR-88': 'UC-29',
      'FR-89': 'UC-26', 'FR-90': 'UC-30', 'FR-91': 'UC-28', 'FR-92': 'UC-28', 'FR-93': 'UC-28', 'FR-94': 'UC-28', 'FR-95': 'UC-27', 'FR-96': 'UC-27',
      'FR-97': 'UC-28', 'FR-98': 'UC-30', 'FR-99': 'UC-12', 'FR-100': 'UC-08', 'FR-101': 'UC-08'}
CHANGED = {'/vocabulary', '/items/:id', '/outfits', '/outfits/wear', '/intelligence/health', '/notifications', '/packing/plan · /packing/trips', '/chat', '/profile'}
NEW = PatternFill('solid', fgColor='E3F2E6')
wrap = Alignment(wrap_text=True, vertical='top')

wb = openpyxl.load_workbook(f'{SRC}/05_Testing/Closet_Traceability_Matrix.xlsx')
ws = wb['Traceability v1.4']; ws.title = 'Traceability v1.5'
for f in FRS:
    rel = 'S2 · changed in v1.5' if f[0] in ('FR-100', 'FR-101') else 'S3 · built early (v1.5 preview); story proposed, not yet in Jira'
    ws.append([f[0], 'Proactive wardrobe intelligence (SRS 4.21)' if f[0] not in ('FR-100', 'FR-101') else 'Recommendations / saved looks', f[1], UC.get(f[0], '—'), f[2], f[5], f[4], f[7], rel])
    for cell in ws[ws.max_row]: cell.fill = NEW; cell.alignment = wrap
rm = wb['Read me']
rm.append(['v1.5 (Milestone 2 final): rows FR-81 – FR-101 added (green). Use-case IDs UC-23 – UC-30 are in SRS v1.2 R4.4; stories US-94 – US-111 are proposed (Proposed_Future_Backlog.md), not yet in Jira.'])
rm.append(['Every test ID in column H exists in the final run (497/497 passing; see Test Plan and Report v1.2, Appendix A). “E2E sNN” = checkpoint NN of e2e/stateful_journey.py.'])
wb.save(f'{OUT}/Closet_Traceability_Matrix_v1.5.xlsx')

wb = openpyxl.load_workbook(f'{SRC}/04_Design/Closet_API_Table.xlsx')
ws = wb['API v1.4']; ws.title = 'API v1.5'
n = ws.max_row - 1
ws2 = wb.create_sheet('v1.5 changes', 1)
ws2.append(['#', 'Method', 'Path (base /api/v1)', 'Purpose', 'Access', 'Request', 'Response', 'Errors', 'New / changed'])
for c in ws2[1]: c.font = Font(bold=True)
existing = {(str(r[1]), str(r[2])) for r in ws.iter_rows(min_row=2, values_only=True)}
for i, a in enumerate(API, 1):
    changed = a[1] in CHANGED
    ws2.append([i, a[0], a[1], a[2], a[3], a[4], a[5], a[6], 'changed' if changed else 'new'])
    if not changed:
        n += 1; ws.append([n, a[0], a[1], a[3].replace('JWT', 'Auth'), 'Proactive intelligence / v1.5', 'S3 preview (v1.5)'])
        for cell in ws[ws.max_row]: cell.fill = NEW
for col, w in zip('ABCDEFGHI', [4, 8, 34, 48, 16, 40, 40, 16, 10]): ws2.column_dimensions[col].width = w
for row in ws2.iter_rows(min_row=2):
    for cell in row: cell.alignment = wrap
wb['Notes'].append(['v1.5: new routes appended (green) to “API v1.5”; sheet “v1.5 changes” lists every new or changed route with request, response and errors (same as SDD v1.2 R4.9).'])
wb.save(f'{OUT}/Closet_API_Table_v1.5.xlsx')
print('ok', n)
